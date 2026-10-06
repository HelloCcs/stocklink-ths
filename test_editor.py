import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
from openpyxl import load_workbook
from main import MainWindow


class EditorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication(['test', '-platform', 'offscreen'])

    def test_edit_delete_undo_save(self):
        window = MainWindow()
        with patch.object(QFileDialog, 'getOpenFileName', return_value=('回测.xlsx', '')):
            window.import_file()
        count = window.table.rowCount()
        self.assertGreater(count, 0)
        window.table.item(0, 4).setText('123.45')
        window.undo()
        self.assertNotEqual(window.table.item(0, 4).text(), '123.45')
        window.redo()
        self.assertEqual(window.table.item(0, 4).text(), '123.45')
        window.table.selectRow(1)
        window.delete_row()
        self.assertEqual(window.table.rowCount(), count - 1)
        window.undo()
        self.assertEqual(window.table.rowCount(), count)
        with tempfile.TemporaryDirectory() as folder:
            target = str(Path(folder) / 'saved.xlsx')
            with patch.object(QFileDialog, 'getSaveFileName', return_value=(target, '')):
                self.assertTrue(window.save_as())
            book = load_workbook(target)
            self.assertEqual(book.active['E3'].value, 123.45)
            self.assertEqual(book.active.max_row, count + 2)
            book.close()
        window.table.item(0, 4).setText('999')
        with patch.object(QMessageBox, 'question', return_value=QMessageBox.Cancel):
            self.assertFalse(window.confirm_discard())
        window.saved_state = window.snapshot()
        window.close()

    def test_paste_undo_and_empty_link(self):
        window = MainWindow()
        with patch.object(QFileDialog, 'getOpenFileName', return_value=('回测.xlsx', '')):
            window.import_file()
        original = window.snapshot()
        window.table.setCurrentCell(0, 0)
        QApplication.clipboard().setText('000001\t平安银行\t2026-01-01\t2026-02-01')
        window.paste_cells()
        self.assertEqual(window.table.item(0, 0).text(), '000001')
        window.undo()
        self.assertEqual(window.snapshot(), original)
        window.add_row()
        with patch.object(QMessageBox, 'warning') as warning:
            window.open_stock(window.table.rowCount() - 1, 0)
            warning.assert_called_once()
        window.saved_state = window.snapshot()
        window.close()

    def test_sort_preserves_rows_and_save_order(self):
        window = MainWindow()
        with patch.object(QFileDialog, 'getOpenFileName', return_value=('回测.xlsx', '')):
            window.import_file()
        original = window.snapshot()
        window.sort_rows(4)
        sorted_state = window.snapshot()
        prices = [float(row[4]) for row in sorted_state[0] if row[4]]
        self.assertEqual(prices, sorted(prices))
        for row, origin in zip(*sorted_state):
            self.assertEqual(row, original[0][origin])
        window.undo()
        self.assertEqual(window.snapshot(), original)
        window.redo()
        self.assertEqual(window.snapshot(), sorted_state)
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / 'sorted.xlsx'
            self.assertTrue(window.write_document(target))
            book = load_workbook(target)
            for offset, row in enumerate(sorted_state[0], 3):
                self.assertEqual(book.active.cell(offset, 2).value, row[1] or None)
            book.close()
            window.table.item(0, 4).setText('42')
            self.assertTrue(window.save())
            book = load_workbook(target)
            self.assertEqual(book.active['E3'].value, 42)
            book.close()
        window.close()


if __name__ == '__main__': unittest.main()
