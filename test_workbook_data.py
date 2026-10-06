import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from workbook_data import WorkbookData, parse_date, stock_code


class WorkbookTests(unittest.TestCase):
    def test_sample_roundtrip(self):
        source = Path(__file__).with_name('回测.xlsx')
        doc = WorkbookData(source)
        original = source.read_bytes()
        self.assertEqual(doc.header_row, 2)
        self.assertEqual(doc.display(0, 2), '2026-08-28')
        doc.book.create_sheet('保留页')['A1'] = '=1+2'
        doc.set_text(0, 0, '000001')
        doc.set_text(0, 4, '12.34')
        doc.set_text(0, 6, '5.25%')
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'saved.xlsx'
            doc.save(path)
            result = load_workbook(path)
            self.assertEqual(result.active['A3'].value, '000001')
            self.assertEqual(result.active['E3'].value, 12.34)
            self.assertEqual(result.active['G3'].value, .0525)
            self.assertEqual(result['保留页']['A1'].value, '=1+2')
            self.assertEqual(result.active['G3'].number_format, '0.00%')
            result.close()
        self.assertEqual(source.read_bytes(), original)

    def test_parsing(self):
        self.assertEqual(stock_code("'000001"), '000001')
        self.assertEqual(parse_date(20260828).isoformat(), '2026-08-28')
        with self.assertRaises(ValueError): parse_date('20260230')


if __name__ == '__main__':
    unittest.main()
