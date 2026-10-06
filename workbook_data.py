"""Workbook editing with field-aware display and atomic saving."""
import os
import copy
from io import BytesIO
import re
import math
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

HEADERS = ['股票代码', '股票简称', '起始时间', '终止时间', '起始价', '终止价', '区间涨跌幅']


def normalize_header(value):
    text = re.sub(r'\s+', '', str(value or ''))
    return {'股票名称': '股票简称', '开始时间': '起始时间', '开始日期': '起始时间',
            '起始日期': '起始时间', '结束时间': '终止时间', '结束日期': '终止时间',
            '终止日期': '终止时间', '起始价格': '起始价', '终止价格': '终止价'}.get(text, text)


def parse_number(text, percentage=False):
    text = str(text).strip()
    if text.endswith('%') and not percentage:
        raise ValueError('价格不能使用百分数')
    value = float(text.rstrip('%')) / (100 if text.endswith('%') else 1)
    if not math.isfinite(value):
        raise ValueError('数值不能为 NaN 或无穷大')
    return value


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if re.fullmatch(r'\d{8}\.0', text):
        text = text[:-2]
    for fmt in ('%Y%m%d', '%Y-%m-%d', '%Y/%m/%d', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f'无法识别日期：{value}')


def stock_code(value):
    text = str(value or '').strip().lstrip("'")
    if text.endswith('.0'):
        text = text[:-2]
    if text.isdigit() and len(text) <= 6:
        return text.zfill(6)
    raise ValueError(f'需要六位股票代码：{value}')


class WorkbookData:
    def clone(self):
        buffer = BytesIO()
        self.book.save(buffer)
        buffer.seek(0)
        result = copy.copy(self)
        result.book = load_workbook(buffer, data_only=False)
        result.sheet = result.book[self.sheet.title]
        return result

    def __init__(self, path, sheet=None):
        self.path = Path(path)
        self.book = load_workbook(path, data_only=False)
        self.sheet = self.book[sheet] if sheet else self.book.active
        for row in self.sheet.iter_rows(max_row=min(30, self.sheet.max_row)):
            names = [normalize_header(c.value) for c in row]
            if all(h in names for h in HEADERS[:4]):
                self.header_row = row[0].row
                self.headers = names
                break
        else:
            self.book.close()
            raise ValueError('前30行未找到股票代码、股票简称、起始时间、终止时间表头')
        if len(set(self.headers)) != len(self.headers):
            raise ValueError('表头重复或存在多个空列，请先修正表头')

    @property
    def row_count(self):
        return max(0, self.sheet.max_row - self.header_row)

    def cell(self, row, col):
        return self.sheet.cell(self.header_row + row + 1, col + 1)

    def display(self, row, col):
        cell = self.cell(row, col)
        value = cell.value
        if value is None:
            return ''
        if cell.data_type == 'f':
            return value
        field = self.headers[col]
        try:
            if field == '股票代码':
                return stock_code(value)
            if field in HEADERS[2:4]:
                return parse_date(value).isoformat()
            if field == '区间涨跌幅' and isinstance(value, (float, int)):
                return f'{value:.2%}'
        except ValueError:
            pass
        return str(value)

    def set_text(self, row, col, text):
        cell = self.cell(row, col)
        if text == self.display(row, col):
            return
        field = self.headers[col]
        text = text.strip()
        value = text or None
        if text and not text.startswith('='):
            if field == '股票代码':
                try:
                    value = stock_code(text)
                except ValueError:
                    value = text
                cell.number_format = '@'
            elif field in HEADERS[2:4]:
                value = parse_date(text)
                cell.number_format = 'yyyy-mm-dd'
            elif field in HEADERS[4:]:
                value = parse_number(text, percentage=field == '区间涨跌幅')
                if field == '区间涨跌幅':
                    cell.number_format = '0.00%'
        cell.value = value

    def save(self, path):
        target = Path(path).resolve()
        fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=target.parent)
        os.close(fd)
        try:
            self.book.save(temporary)
            os.replace(temporary, target)
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)
        self.path = target
