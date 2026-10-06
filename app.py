from __future__ import annotations

import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QApplication, QFileDialog, QLabel, QMainWindow, QMessageBox, QPushButton,
    QTableWidget, QTableWidgetItem, QToolBar, QVBoxLayout, QWidget
)

from adapters import AdapterResult, create_adapter


REQUIRED = ["股票代码", "股票简称", "起始时间", "终止时间"]
OPTIONAL = ["起始价", "终止价", "区间涨跌幅"]


def clean_code(value) -> str:
    if value is None:
        return ""
    return str(value).strip().lstrip("'")


def parse_date(value) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace("/", "-")
    for fmt in ("%Y%m%d", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def display_value(value, column: str) -> str:
    if value is None:
        return ""
    if column in ("起始时间", "终止时间"):
        parsed = parse_date(value)
        return parsed.isoformat() if parsed else str(value)
    return str(value)


@dataclass
class SheetData:
    headers: list[str]
    rows: list[list[object]]


class StockApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("股票表格与同花顺联动工具")
        self.resize(1120, 680)
        self.current_path: Optional[Path] = None
        self.modified = False
        self.adapter = create_adapter()

        self.table = QTableWidget(0, 0)
        self.table.setAlternatingRowColors(True)
        self.table.setSortingEnabled(True)
        self.table.cellDoubleClicked.connect(self.on_double_click)
        self.table.itemChanged.connect(lambda _item: self.set_modified(True))
        self.setCentralWidget(self.table)

        toolbar = QToolBar()
        self.addToolBar(toolbar)
        for label, callback in (
            ("导入表格", self.import_file),
            ("另存为", self.save_as),
            ("校验", self.validate_rows),
            ("新增行", self.add_row),
            ("删除行", self.delete_rows),
            ("联动测试", self.test_adapter),
        ):
            button = QPushButton(label)
            button.clicked.connect(callback)
            toolbar.addWidget(button)
        self.info = QLabel("未导入文件 | 同花顺：" + self.adapter.name)
        self.statusBar().addPermanentWidget(self.info)

    def set_modified(self, value=True):
        self.modified = value
        if self.current_path:
            self.setWindowTitle(("*" if value else "") + f"股票表格与同花顺联动工具 - {self.current_path.name}")

    def import_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "选择 Excel 文件", "", "Excel 文件 (*.xlsx)")
        if not path:
            return
        try:
            workbook = load_workbook(path, data_only=False)
            sheet = workbook.active
            header_row = None
            for row_idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
                values = [str(v).strip() if v is not None else "" for v in row]
                if all(field in values for field in REQUIRED):
                    header_row = row_idx
                    break
            if header_row is None:
                raise ValueError("未找到包含股票代码、股票简称、起始时间、终止时间的表头行")
            headers = [str(v).strip() if v is not None else f"未命名列{i+1}" for i, v in enumerate(next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))]
            rows = [[cell for cell in row] for row in sheet.iter_rows(min_row=header_row + 1, max_col=len(headers), values_only=True)]
            while rows and all(v is None for v in rows[-1]):
                rows.pop()
            self.load_data(SheetData(headers, rows))
            self.current_path = Path(path)
            self.set_modified(False)
            self.info.setText(f"{self.current_path.name} | {len(rows)} 行 | 同花顺：{self.adapter.name}")
        except Exception as exc:
            QMessageBox.critical(self, "导入失败", str(exc))

    def load_data(self, data: SheetData):
        self.table.blockSignals(True)
        self.table.setSortingEnabled(False)
        self.table.clear()
        self.table.setColumnCount(len(data.headers))
        self.table.setHorizontalHeaderLabels(data.headers)
        self.table.setRowCount(len(data.rows))
        for r, row in enumerate(data.rows):
            for c, value in enumerate(row):
                self.table.setItem(r, c, QTableWidgetItem(display_value(value, data.headers[c])))
        self.table.resizeColumnsToContents()
        self.table.setSortingEnabled(True)
        self.table.blockSignals(False)

    def table_data(self):
        headers = [self.table.horizontalHeaderItem(c).text() for c in range(self.table.columnCount())]
        rows = []
        for r in range(self.table.rowCount()):
            rows.append([self.table.item(r, c).text() if self.table.item(r, c) else "" for c in range(self.table.columnCount())])
        return headers, rows

    def validate_rows(self):
        headers, rows = self.table_data()
        positions = {name: headers.index(name) for name in REQUIRED if name in headers}
        errors = []
        if len(positions) != len(REQUIRED):
            errors.append("缺少必填列：" + ", ".join(set(REQUIRED) - set(positions)))
        for row_num, row in enumerate(rows, 1):
            code = clean_code(row[positions["股票代码"]]) if "股票代码" in positions else ""
            start = parse_date(row[positions["起始时间"]]) if "起始时间" in positions else None
            end = parse_date(row[positions["终止时间"]]) if "终止时间" in positions else None
            if not code:
                errors.append(f"第 {row_num} 行：股票代码为空")
            if not start or not end:
                errors.append(f"第 {row_num} 行：起始时间/终止时间无法解析")
            elif start > end:
                errors.append(f"第 {row_num} 行：起始时间晚于终止时间")
        if errors:
            QMessageBox.warning(self, "校验结果", "\n".join(errors[:30]) + ("\n……" if len(errors) > 30 else ""))
            return False
        QMessageBox.information(self, "校验结果", f"校验通过，共 {len(rows)} 行。")
        return True

    def save_as(self):
        if self.table.columnCount() == 0:
            QMessageBox.information(self, "提示", "请先导入表格")
            return
        path, _ = QFileDialog.getSaveFileName(self, "另存为", "编辑后.xlsx", "Excel 文件 (*.xlsx)")
        if not path:
            return
        headers, rows = self.table_data()
        wb = Workbook()
        ws = wb.active
        ws.title = "股票数据"
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = cell.font.copy(bold=True)
        for idx, header in enumerate(headers, 1):
            if header == "区间涨跌幅":
                for row in range(2, ws.max_row + 1):
                    try:
                        ws.cell(row, idx).value = float(ws.cell(row, idx).value)
                        ws.cell(row, idx).number_format = "0.00%"
                    except (TypeError, ValueError):
                        pass
            if header == "股票代码":
                for row in range(2, ws.max_row + 1):
                    ws.cell(row, idx).number_format = "@"
                    ws.cell(row, idx).value = clean_code(ws.cell(row, idx).value)
        wb.save(path)
        self.current_path = Path(path)
        self.set_modified(False)
        self.info.setText(f"已保存：{self.current_path.name} | {len(rows)} 行 | 同花顺：{self.adapter.name}")

    def add_row(self):
        if self.table.columnCount() == 0:
            return
        self.table.insertRow(self.table.rowCount())
        self.set_modified(True)

    def delete_rows(self):
        selected = sorted({idx.row() for idx in self.table.selectedIndexes()}, reverse=True)
        for row in selected:
            self.table.removeRow(row)
        if selected:
            self.set_modified(True)

    def on_double_click(self, row, column):
        header = self.table.horizontalHeaderItem(column).text()
        if header not in ("股票代码", "股票简称"):
            return
        headers, values = self.table_data()
        mapping = dict(zip(headers, values[row]))
        result: AdapterResult = self.adapter.open_stock(
            clean_code(mapping.get("股票代码")),
            mapping.get("股票简称", ""),
            parse_date(mapping.get("起始时间")),
            parse_date(mapping.get("终止时间")),
        )
        QMessageBox.information(self, "同花顺联动", result.message) if result.ok else QMessageBox.warning(self, "同花顺联动", result.message)

    def test_adapter(self):
        result = self.adapter.open_stock("000001", "平安银行", date.today(), date.today())
        QMessageBox.information(self, "联动测试", result.message)


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = StockApp()
    window.show()
    sys.exit(app.exec())
