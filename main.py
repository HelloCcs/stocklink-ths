import sys

if __name__ == '__main__' and '--vision-self-test' in sys.argv:
    from vision_probe import main as run_vision_test
    run_vision_test(sys.argv[sys.argv.index('--vision-self-test') + 1])
    raise SystemExit(0)

if __name__ == '__main__' and '--ocr-worker' in sys.argv:
    # Run before Qt initializes an STA apartment; WinRT OCR needs its own loop.
    from local_ocr import main as run_ocr
    run_ocr(sys.argv[sys.argv.index('--ocr-worker') + 1:])
    raise SystemExit(0)

if __name__ == '__main__' and '--mark-worker' in sys.argv:
    from chart_tracking import main as run_marking
    run_marking(sys.argv[sys.argv.index('--mark-worker') + 1:])
    raise SystemExit(0)

import copy
import ctypes
import re
from pathlib import Path
from datetime import datetime, date
import time
from workbook_data import WorkbookData, parse_date, stock_code, parse_number
from runtime_log import configure_log, log_path
from app_config import load_config, save_config

from openpyxl import load_workbook, Workbook
from PySide6.QtCore import Qt, QThread, Signal, QTimer, QRect, QObject, QAbstractNativeEventFilter
from PySide6.QtGui import QAction, QKeySequence
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QPushButton,
    QFileDialog, QMessageBox, QTableWidget, QTableWidgetItem, QLabel, QAbstractItemView, QInputDialog, QCheckBox,
    QDialog, QFormLayout, QDialogButtonBox, QComboBox, QSpinBox
)

HEADERS = ['股票代码', '股票简称', '起始时间', '终止时间', '起始价', '终止价', '区间涨跌幅']

class LinkJob(QThread):
    result = Signal(object)

    def __init__(self, values, parent):
        super().__init__(parent)
        self.values = values
        self.target = parent.target_window
        self.navigate = True
        self.timeout = parent.config.get('link_timeout', 180)
        self.direction = parent.config.get('date_direction', 'nearest')

    def run(self):
        from adapters import TongHuaShunAdapter
        try:
            if self.navigate:
                result = TongHuaShunAdapter(self.target, self.isInterruptionRequested,
                                            self.timeout, self.direction).open_stock(*self.values)
            else:
                from marking_client import locate_markers
                from adapters import AdapterResult
                code, name, start, end = self.values
                markers = locate_markers(self.target, code, start, end, timeout=self.timeout,
                                         direction=self.direction, cancelled=self.isInterruptionRequested)
                result = AdapterResult(True, f'已重新定位 {code} 的日期标记', markers)
            self.result.emit(result)
        except Exception as exc:
            from adapters import AdapterResult
            self.result.emit(AdapterResult(False, f'联动失败：{exc}'))


class DailyJumpJob(QThread):
    result = Signal(object)

    def __init__(self, code, parent):
        super().__init__(parent)
        self.code = code
        self.target = copy.copy(parent.target_window) if parent is not None else None

    def run(self):
        from adapters import TongHuaShunAdapter
        try:
            wait_for_jump_key_release()
            self.result.emit(TongHuaShunAdapter(target=self.target, timeout=30).open_stock_daily(self.code))
        except Exception as exc:
            from adapters import AdapterResult
            self.result.emit(AdapterResult(False, f'快捷跳转失败：{exc}'))


def wait_for_jump_key_release(user=None, clock=time.monotonic, sleep=time.sleep):
    """Never send Escape/digits while the triggering Ctrl/Alt is held."""
    user = user if user is not None else ctypes.WinDLL('user32', use_last_error=True)
    user.GetAsyncKeyState.argtypes = [ctypes.c_int]
    user.GetAsyncKeyState.restype = ctypes.c_short
    deadline = clock() + 3
    for _ in range(150):
        if clock() >= deadline:
            break
        if not any(user.GetAsyncKeyState(key) & 0x8000
                   for key in (16, 17, 18, 91, 92, ord('Z'))):
            return
        sleep(.02)
    raise RuntimeError('快捷键仍被按住，已取消发送；请松开 Ctrl、Alt 后重试')


class GlobalJumpHotkey(QAbstractNativeEventFilter):
    """Register Ctrl+Alt+Z and forward the clipboard code to MainWindow."""
    HOTKEY_ID = 0x534C
    WM_HOTKEY = 0x0312
    MOD_ALT, MOD_CONTROL = 0x0001, 0x0002
    VK_Z = 0x5A

    class MSG(ctypes.Structure):
        _fields_ = [('hwnd', ctypes.c_void_p), ('message', ctypes.c_uint),
                    ('wParam', ctypes.c_size_t), ('lParam', ctypes.c_ssize_t),
                    ('time', ctypes.c_uint), ('pt_x', ctypes.c_long), ('pt_y', ctypes.c_long)]

    def __init__(self, window, callback):
        super().__init__()
        self.window = window
        self.callback = callback
        self.user32 = None
        self.registered = False
        self.shortcut = 'Ctrl+Alt+Z'

    def register(self):
        if sys.platform != 'win32':
            return False
        self.user32 = ctypes.WinDLL('user32', use_last_error=True)
        from ctypes import wintypes
        self.user32.RegisterHotKey.argtypes = [wintypes.HWND, ctypes.c_int, wintypes.UINT, wintypes.UINT]
        self.user32.RegisterHotKey.restype = wintypes.BOOL
        self.user32.UnregisterHotKey.argtypes = [wintypes.HWND, ctypes.c_int]
        self.user32.UnregisterHotKey.restype = wintypes.BOOL
        hwnd = int(self.window.winId())
        self.registered = bool(self.user32.RegisterHotKey(hwnd, self.HOTKEY_ID,
                                                          self.MOD_CONTROL | self.MOD_ALT | 0x4000, self.VK_Z))
        if self.registered:
            QApplication.instance().installNativeEventFilter(self)
        return self.registered

    def unregister(self):
        if self.registered and self.user32 is not None:
            self.user32.UnregisterHotKey(int(self.window.winId()), self.HOTKEY_ID)
            QApplication.instance().removeNativeEventFilter(self)
            self.registered = False

    def nativeEventFilter(self, event_type, message):
        if self.registered and message is not None and int(message) != 0:
            msg = ctypes.cast(int(message), ctypes.POINTER(self.MSG)).contents
            if msg.message == self.WM_HOTKEY and msg.wParam == self.HOTKEY_ID:
                self.callback()
                return True, 0
        return False, 0


def clipboard_stock_code(text):
    """Return one unambiguous six-digit code from clipboard text."""
    candidates = re.findall(r'(?<!\d)(\d{6})(?!\d)', str(text or ''))
    unique = list(dict.fromkeys(candidates))
    if len(unique) != 1:
        raise ValueError('剪贴板中必须包含唯一的六位股票代码')
    return unique[0]

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.logger = configure_log()
        self.path = None
        self.document = None
        self.history = []
        self.future = []
        self.saved_state = None
        self.last_state = None
        self.link_job = None
        self.pending_link = None
        self.target_window = None
        self.marker_payload = None
        self.marker_watch = None
        self.marker_retry_after = 0
        self.marker_recalibrating = False
        self.marker_recalibration_attempts = 0
        self.marker_recalibration_max_attempts = 3
        self.marker_recalibration_deadline = 0
        self.marker_recalibration_error = None
        from chart_overlay import DateMarkerOverlay
        self.marker_overlay = DateMarkerOverlay()
        self.marker_timer = QTimer(self)
        self.marker_timer.timeout.connect(self._refresh_marker_window)
        self.marker_timer.start(500)
        self.config = load_config()
        self.sort_column = None
        self.sort_descending = False
        self.headers = HEADERS[:]
        self.table = QTableWidget(0, len(self.headers))
        self.table.setHorizontalHeaderLabels(self.headers)
        self.table.setEditTriggers(QAbstractItemView.EditKeyPressed)
        self.table.setToolTip('单击代码/简称联动；选中单元格按 F2 编辑。可用 Ctrl+C / Ctrl+V 复制粘贴。')
        self.table.cellClicked.connect(self.open_stock)
        self.table.itemChanged.connect(self.record_change)
        self.table.horizontalHeader().sectionClicked.connect(self.sort_rows)
        self.status = QLabel('状态：离线模式，未检测到同花顺')
        saved_target = self.config.get('target_window')
        if isinstance(saved_target, dict) and saved_target.get('hwnd') and saved_target.get('pid'):
            self.target_window = saved_target
            self.status.setText(f"已恢复目标窗口：HWND {saved_target['hwnd']} / PID {saved_target['pid']}")
        bar = QHBoxLayout()
        self.real_mode = QCheckBox('启用真实同花顺联动')
        self.real_mode.toggled.connect(self.mode_changed)
        # Real linkage is the production default; users can uncheck for offline preview.
        self.real_mode.setChecked(True)
        bar.addWidget(self.real_mode)
        choose = QPushButton('选择看盘窗口')
        choose.clicked.connect(self.choose_target)
        bar.addWidget(choose)
        for text, slot in [('导入表格', self.import_file), ('另存为', self.save_as), ('校验', self.validate), ('添加行', self.add_row), ('删除行', self.delete_row)]:
            b = QPushButton(text); b.clicked.connect(slot); bar.addWidget(b)
        root = QVBoxLayout(); root.addLayout(bar); root.addWidget(self.table); root.addWidget(self.status)
        edit_bar = QHBoxLayout()
        for label, callback in [('保存', self.save), ('撤销', self.undo), ('重做', self.redo), ('设置', self.show_settings), ('停止联动', self.cancel_link), ('复制诊断信息', self.copy_diagnostics)]:
            button = QPushButton(label)
            button.clicked.connect(callback)
            edit_bar.addWidget(button)
        root.insertLayout(1, edit_bar)
        w = QWidget(); w.setLayout(root); self.setCentralWidget(w)
        self.setWindowTitle('股票表格与同花顺联动工具'); self.resize(1000, 650)
        self.daily_jump_job = None
        self.global_jump_hotkey = GlobalJumpHotkey(self, self.handle_global_jump)
        QTimer.singleShot(0, self._register_global_jump_hotkey)
        for label, shortcut, callback in [('保存', 'Ctrl+S', self.save), ('撤销', 'Ctrl+Z', self.undo), ('重做', 'Ctrl+Y', self.redo), ('复制', 'Ctrl+C', self.copy_cells), ('粘贴', 'Ctrl+V', self.paste_cells), ('另存为', 'Ctrl+Shift+S', self.save_as)]:
            action = QAction(label, self)
            action.setShortcut(QKeySequence(shortcut))
            action.triggered.connect(callback)
            self.addAction(action)

    def _register_global_jump_hotkey(self):
        if self.global_jump_hotkey.register():
            self.status.setText(f'快捷键已就绪：WPS 中 Ctrl+C 复制代码，再按 {self.global_jump_hotkey.shortcut}')
            self.logger.info('global_hotkey_registered %s', self.global_jump_hotkey.shortcut)
        elif sys.platform == 'win32':
            self.logger.warning('global_hotkey_register_failed error=%s', ctypes.get_last_error())
            self.status.setText('Ctrl+Alt+Z 注册失败：请关闭其他 StockLink 或占用该快捷键的程序后重启')

    def handle_global_jump(self):
        self.logger.info('global_hotkey_received')
        if self.link_job is not None or self.marker_watch is not None:
            self.status.setText('原表格联动仍在执行，请停止后再使用快捷键')
            return
        if self.daily_jump_job is not None:
            self.status.setText('上一次快捷跳转仍在执行')
            return
        try:
            code = clipboard_stock_code(QApplication.clipboard().text())
        except ValueError as exc:
            self.status.setText(str(exc))
            return
        self.status.setText(f'正在跳转 {code} 到同花顺日K…')
        job = DailyJumpJob(code, self)
        self.daily_jump_job = job
        job.result.connect(self.global_jump_result)
        job.finished.connect(self.global_jump_finished)
        job.start()

    def global_jump_result(self, result):
        self.logger.info('global_jump_result %s', getattr(result, 'message', str(result)))
        self.status.setText(getattr(result, 'message', str(result)))

    def global_jump_finished(self):
        job, self.daily_jump_job = self.daily_jump_job, None
        if job is not None:
            job.deleteLater()

    def show_settings(self):
        dialog = QDialog(self)
        dialog.setWindowTitle('联动设置')
        form = QFormLayout(dialog)
        direction = QComboBox()
        for label, value in [('最近交易日（等距取前一日）', 'nearest'), ('向前取交易日', 'previous'), ('向后取交易日', 'next')]:
            direction.addItem(label, value)
        direction.setCurrentIndex(max(0, direction.findData(self.config.get('date_direction', 'nearest'))))
        timeout = QSpinBox()
        timeout.setRange(30, 900)
        timeout.setSuffix(' 秒')
        timeout.setValue(self.config.get('link_timeout', 180))
        form.addRow('非交易日定位', direction)
        form.addRow('日期定位超时', timeout)
        form.addRow('日期标记', QLabel('透明覆盖层'))
        log = QLabel(log_path(self.logger))
        log.setWordWrap(True)
        log.setTextInteractionFlags(Qt.TextSelectableByMouse)
        form.addRow('日志目录', log)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        form.addRow(buttons)
        if dialog.exec() == QDialog.Accepted:
            candidate = dict(self.config, date_direction=direction.currentData(), link_timeout=timeout.value())
            if not self.persist_config(candidate):
                return
            self.cancel_link()
            self.status.setText('联动设置已更新，请重新点击股票定位日期')

    def persist_config(self, candidate):
        try:
            save_config(candidate)
        except OSError:
            self.logger.exception('config_save_failed')
            QMessageBox.warning(self, '设置未保存', '无法写入本机配置目录，原设置保持不变。请检查目录权限或磁盘空间后重试。')
            return False
        self.config = candidate
        return True

    def copy_diagnostics(self):
        import json
        from diagnose_ths import collect_report
        report = collect_report()
        report['selected_target'] = self.target_window
        report['last_status'] = self.status.text()
        report['date_marking_verified'] = False
        report['log_file'] = log_path(self.logger)
        QApplication.clipboard().setText(json.dumps(report, ensure_ascii=False, indent=2))
        self.status.setText('诊断信息已复制')

    def sort_rows(self, column):
        if self.document is None: return
        self.sort_descending = not self.sort_descending if self.sort_column == column else False
        self.sort_column = column
        before = self.snapshot()
        rows, origins = before
        field = self.headers[column]
        def key(index):
            text = rows[index][column].strip()
            try:
                if field in HEADERS[2:4]: return (0, parse_date(text))
                if field in HEADERS[4:]: return (0, parse_number(text, field == '区间涨跌幅'))
                if field == '股票代码': return (0, stock_code(text))
            except ValueError:
                return (1, text)
            return (0, text.casefold())
        indices = sorted(range(len(rows)), key=key, reverse=self.sort_descending)
        after = ([rows[i] for i in indices], [origins[i] for i in indices])
        if after != before:
            self.history.append(before)
            self.future.clear()
            self.restore(after)
        self.table.horizontalHeader().setSortIndicatorShown(True)
        self.table.horizontalHeader().setSortIndicator(column, Qt.DescendingOrder if self.sort_descending else Qt.AscendingOrder)

    def choose_target(self):
        from diagnose_ths import collect_windows
        report = collect_windows()
        windows = report.get('windows', [])
        if not windows:
            QMessageBox.warning(self, '目标窗口', report.get('status', '未找到同花顺窗口'))
            return
        labels = [f"{i+1}. {w['title'] or w['class']} | PID {w['pid']} | HWND {w['hwnd']}" for i, w in enumerate(windows)]
        label, ok = QInputDialog.getItem(self, '选择看盘窗口', '选择需要联动的屏幕/窗口：', labels, 0, False)
        if ok:
            target = windows[labels.index(label)]
            if not self.persist_config(dict(self.config, target_window=target)):
                return
            self.cancel_link()
            self.target_window = target
            self.logger.info('target_selected hwnd=%s pid=%s', self.target_window['hwnd'], self.target_window['pid'])
            self.status.setText('目标窗口：' + label)

    def mode_changed(self, enabled):
        if not enabled:
            self.pending_link = None
            self.marker_payload = None
            self.marker_overlay.clear()
            if self.link_job is not None:
                self.link_job.requestInterruption()
            self.status.setText('已取消待执行联动；正在执行的请求结束后停止')

    def cancel_link(self):
        self.pending_link = None
        self.marker_payload = None
        self.marker_overlay.clear()
        if self.link_job is not None:
            self.link_job.requestInterruption()
        self.status.setText('正在停止联动，已隐藏日期标记')

    def submit_link(self, values):
        request = (values, copy.copy(self.target_window))
        if self.link_job is not None:
            self.pending_link = request
            self.status.setText(f'已排队：{values[0]}（保留最后一次选择）')
            return
        self.start_link(request)

    def start_link(self, request):
        self.marker_payload = None
        self.marker_overlay.clear()
        self.marker_recalibrating = False
        self.marker_recalibration_attempts = 0
        self.marker_recalibration_deadline = 0
        self.marker_recalibration_error = None
        values, target = request
        job = LinkJob(values, self)
        job.target = target
        self.link_job = job
        job.result.connect(self.link_result)
        job.finished.connect(self.link_finished)
        self.status.setText(f'正在向同花顺发送 {values[0]}…')
        self.logger.info('link_started code=%s start=%s end=%s', values[0], values[2], values[3])
        job.start()

    def link_result(self, result):
        message = getattr(result, 'message', str(result))
        self.status.setText(message)
        self.logger.info('link_result %s', message.replace('\n', ' | '))
        if (not self.real_mode.isChecked() or self.pending_link is not None
                or (self.link_job is not None and self.link_job.isInterruptionRequested())):
            return
        payload = getattr(result, 'markers', None)
        if getattr(result, 'ok', False) and payload:
            self.marker_payload = payload
            if self.marker_recalibrating:
                self.marker_recalibrating = False
                self.marker_recalibration_error = None
            # Validate again after worker exit, before showing a potentially
            # stale result on the user's current screen.
            self._refresh_marker_window(force=True)
        else:
            self.marker_overlay.clear()
            if self.marker_recalibrating:
                self.marker_recalibration_error = message
                self.marker_recalibrating = False
                self.marker_payload = None
                self.marker_retry_after = float('inf')
                self.status.setText(f'日期标记自动重校准失败：{message}')
            else:
                self.marker_retry_after = time.monotonic() + 10

    def _refresh_marker_window(self, force=False):
        if (not self.marker_payload or self.marker_watch is not None
                or (self.link_job is not None and not force)
                or not self.real_mode.isChecked()):
            return
        from marker_jobs import MarkerWatchJob
        job = MarkerWatchJob(self.marker_payload, self)
        self.marker_watch = job
        job.result.connect(self.marker_checked)
        job.finished.connect(self.marker_watch_finished)
        job.start()

    def marker_watch_finished(self):
        job, self.marker_watch = self.marker_watch, None
        if job is not None:
            job.deleteLater()

    def marker_checked(self, observation):
        payload, state, detail = observation
        if payload is not self.marker_payload:
            return
        if state == 'valid':
            if not self.marker_overlay.isVisible():
                try:
                    self.marker_overlay.present_verified(payload)
                except Exception as exc:
                    self.status.setText(str(exc))
                    self.marker_payload = None
            return
        self.marker_overlay.clear()
        if state == 'closed':
            self.marker_payload = None
        elif (state == 'changed' and self.link_job is None
              and not self.marker_recalibrating
              and self.marker_recalibration_attempts < self.marker_recalibration_max_attempts
              and (self.marker_recalibration_deadline == 0
                   or time.monotonic() < self.marker_recalibration_deadline)
              and time.monotonic() >= self.marker_retry_after):
            values = (payload['code'], '', date.fromisoformat(payload['start']['requested']),
                      date.fromisoformat(payload['end']['requested']))
            now = time.monotonic()
            if self.marker_recalibration_deadline == 0:
                self.marker_recalibration_deadline = now + self.config.get('link_timeout', 180)
            self.marker_recalibration_attempts += 1
            self.marker_recalibrating = True
            job = LinkJob(values, self)
            job.target, job.navigate = payload, False
            self.link_job = job
            job.result.connect(self.link_result)
            job.finished.connect(self.link_finished)
            self.status.setText('图表已变化，正在重新定位日期标记…')
            job.start()
        elif state == 'changed' and self.link_job is None:
            self.marker_overlay.clear()
            self.marker_recalibrating = False
            self.marker_recalibration_error = detail or '自动重校准已达到次数或时间上限'
            self.marker_payload = None
            self.marker_retry_after = float('inf')
            self.status.setText(f'日期标记自动重校准停止：{self.marker_recalibration_error}')

    def link_finished(self):
        completed = self.link_job
        self.link_job = None
        if completed is not None:
            completed.deleteLater()
        request, self.pending_link = self.pending_link, None
        if request and self.real_mode.isChecked():
            self.start_link(request)

    def snapshot(self):
        return ([[self.table.item(r, c).text() if self.table.item(r, c) else ''
                  for c in range(self.table.columnCount())] for r in range(self.table.rowCount())],
                list(getattr(self, 'source_rows', [])))

    def record_change(self, *_):
        if self.document is None: return
        state = self.snapshot()
        if self.last_state is not None and state != self.last_state:
            self.history.append(self.last_state)
            self.future.clear()
        self.last_state = state
        self.status.setText('状态：已修改' if state != self.saved_state else '状态：已保存')

    def restore(self, state):
        rows, self.source_rows = copy.deepcopy(state)
        self.table.blockSignals(True)
        self.table.setRowCount(len(rows))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                self.table.setItem(r, c, QTableWidgetItem(value))
        self.table.blockSignals(False)
        self.last_state = self.snapshot()
        self.table.horizontalHeader().setSortIndicatorShown(False)
        self.status.setText('状态：已修改' if self.last_state != self.saved_state else '状态：已保存')

    def undo(self):
        if self.history:
            self.future.append(self.snapshot())
            self.restore(self.history.pop())

    def redo(self):
        if self.future:
            self.history.append(self.snapshot())
            self.restore(self.future.pop())

    def copy_cells(self):
        ranges = self.table.selectedRanges()
        if not ranges: return
        area = ranges[0]
        rows = self.snapshot()[0]
        QApplication.clipboard().setText('\n'.join('\t'.join(rows[r][area.leftColumn():area.rightColumn()+1]) for r in range(area.topRow(), area.bottomRow()+1)))

    def paste_cells(self):
        if self.document is None or self.table.currentRow() < 0: return
        r, c = self.table.currentRow(), self.table.currentColumn()
        rows = QApplication.clipboard().text().splitlines()
        self.table.blockSignals(True)
        for offset, text in enumerate(rows):
            while r + offset >= self.table.rowCount():
                self.table.insertRow(self.table.rowCount())
                self.source_rows.append(None)
            for delta, value in enumerate(text.split('\t')):
                if c + delta < self.table.columnCount():
                    self.table.setItem(r + offset, c + delta, QTableWidgetItem(value))
        self.table.blockSignals(False)
        self.record_change()

    def confirm_discard(self):
        if self.document is None or self.snapshot() == self.saved_state: return True
        answer = QMessageBox.question(self, '未保存修改', '保存当前修改后继续？', QMessageBox.Save | QMessageBox.Discard | QMessageBox.Cancel, QMessageBox.Cancel)
        if answer == QMessageBox.Save: return self.save_as()
        return answer == QMessageBox.Discard

    def closeEvent(self, event):
        if self.link_job is not None or self.marker_watch is not None or self.daily_jump_job is not None:
            QMessageBox.information(self, '联动执行中', '请等待当前指令执行结束后关闭。')
            event.ignore()
            return
        if self.confirm_discard():
            self.global_jump_hotkey.unregister()
            self.marker_timer.stop()
            self.marker_payload = None
            self.marker_overlay.clear()
            event.accept()
        else:
            event.ignore()

    def import_file(self):
        if not self.confirm_discard(): return
        path, _ = QFileDialog.getOpenFileName(self, '选择 Excel 文件', '', 'Excel 文件 (*.xlsx)')
        if not path: return
        try:
            probe = load_workbook(path, read_only=True)
            sheets = probe.sheetnames
            probe.close()
            sheet = sheets[0]
            if len(sheets) > 1:
                sheet, accepted = QInputDialog.getItem(self, '选择工作表', '导入工作表：', sheets, 0, False)
                if not accepted: return
            document = WorkbookData(path, sheet)
            self.table.blockSignals(True)
            source_headers = document.headers
            self.headers = source_headers
            self.table.clear(); self.table.setColumnCount(len(source_headers)); self.table.setHorizontalHeaderLabels(source_headers)
            self.table.setRowCount(document.row_count)
            for r in range(document.row_count):
                for c in range(len(source_headers)):
                    self.table.setItem(r, c, QTableWidgetItem(document.display(r, c)))
            self.document = document
            self.pending_link = None
            self.sort_column = None
            self.table.horizontalHeader().setSortIndicatorShown(False)
            self.source_rows = list(range(document.row_count))
            self.history.clear(); self.future.clear()
            self.last_state = self.saved_state = self.snapshot()
            self.table.blockSignals(False)
            self.path = Path(path); self.status.setText(f'已导入：{self.path.name}，状态：未修改')
            self.logger.info('import_completed rows=%s columns=%s', document.row_count, len(self.headers))
        except Exception as exc:
            self.logger.exception('import_failed')
            self.table.blockSignals(False)
            QMessageBox.critical(self, '导入失败', str(exc))

    def add_row(self):
        if self.document is None: return
        self.table.insertRow(self.table.rowCount())
        self.source_rows.append(None)
        self.record_change()
    def delete_row(self):
        rows = sorted({i.row() for i in self.table.selectedIndexes()}, reverse=True)
        for row in rows:
            self.source_rows.pop(row)
            self.table.removeRow(row)
        self.record_change()

    def validate(self):
        index = {h: i for i, h in enumerate(self.headers)}
        missing = [h for h in HEADERS[:4] if h not in index]
        errors = []
        if missing: errors.append('缺少字段：' + '、'.join(missing))
        for r in range(self.table.rowCount()):
            def val(name):
                item = self.table.item(r, index.get(name, -1)); return item.text().strip() if item else ''
            try: stock_code(val('股票代码'))
            except ValueError as exc: errors.append(f'第 {r + 1} 行：{exc}')
            try:
                start = parse_date(val('起始时间'))
                end = parse_date(val('终止时间'))
                if start > end: errors.append(f'第 {r + 1} 行起始时间晚于终止时间')
            except ValueError: errors.append(f'第 {r + 1} 行日期格式错误')
            for field in HEADERS[4:]:
                if field in index and val(field) and not val(field).startswith('='):
                    try: parse_number(val(field), percentage=field == '区间涨跌幅')
                    except ValueError as exc: errors.append(f'第 {r + 1} 行 {field}：{exc}')
        QMessageBox.information(self, '校验结果', '校验通过' if not errors else '\n'.join(errors[:30]))

    def open_stock(self, row, column):
        if self.headers[column] not in ('股票代码', '股票简称'): return
        code_col = self.headers.index('股票代码') if '股票代码' in self.headers else -1
        code = self.table.item(row, code_col).text().lstrip("'").strip() if code_col >= 0 and self.table.item(row, code_col) else ''
        values = dict(zip(self.headers, self.snapshot()[0][row]))
        start = values.get('起始时间', '')
        end = values.get('终止时间', '')
        try:
            code = stock_code(code)
            start, end = parse_date(start), parse_date(end)
            if start > end: raise ValueError('起始时间不得晚于终止时间')
        except ValueError as exc:
            QMessageBox.warning(self, '无法联动', str(exc))
            return
        if self.real_mode.isChecked():
            self.submit_link((code, values.get('股票简称', ''), start, end))
            return
        QMessageBox.information(self, '模拟联动', f'离线模式：将打开 {code}\n日K区间：{start} 至 {end}\n目标机安装同花顺后将执行真实联动。')

    def export_marker(self):
        from date_marker import marker_formula
        row = self.table.currentRow()
        if row < 0:
            QMessageBox.information(self, '提示', '请先选择一行')
            return
        values = dict(zip(self.headers, self.snapshot()[0][row]))
        try:
            formula = marker_formula(values.get('起始时间'), values.get('终止时间'))
            path, _ = QFileDialog.getSaveFileName(self, '保存公式验证稿', '日期标记验证稿.txt', '文本 (*.txt)')
            if not path: return
            Path(path).write_text(formula, encoding='utf-8')
            QMessageBox.information(self, '已生成验证稿', '请在目标机同花顺新建主图指标并粘贴验证稿进行编译测试。\n此功能尚未自动导入同花顺，不代表日期标记已经完成。')
        except (ValueError, OSError) as exc:
            QMessageBox.warning(self, '无法生成', str(exc))

    def save(self):
        if self.document is None: return self.save_as()
        return self.write_document(self.path)

    def save_as(self):
        if self.document is None:
            QMessageBox.information(self, '提示', '请先导入表格')
            return
        default = str(self.path.with_name(self.path.stem + '_编辑后.xlsx')) if self.path else '编辑后.xlsx'
        path, _ = QFileDialog.getSaveFileName(self, '另存为', default, 'Excel 文件 (*.xlsx)')
        if not path: return
        if not path.lower().endswith('.xlsx'): path += '.xlsx'
        return self.write_document(path)

    def write_document(self, path):
        try:
            candidate = self.document.clone()
            original = self.document
            for r, source in enumerate(self.source_rows):
                for c in range(self.table.columnCount()):
                    target = candidate.cell(r, c)
                    target.value = original.cell(source, c).value if source is not None else None
                    if source is not None:
                        target._style = copy.copy(original.cell(source, c)._style)
            excess = candidate.row_count - self.table.rowCount()
            if excess > 0:
                candidate.sheet.delete_rows(candidate.header_row + self.table.rowCount() + 1, excess)
            for r in range(self.table.rowCount()):
                for c in range(self.table.columnCount()):
                    item = self.table.item(r, c)
                    candidate.set_text(r, c, item.text() if item else '')
            candidate.save(path)
            self.path = Path(path)
            self.status.setText(f'已保存：{self.path.name}')
            self.logger.info('save_completed rows=%s', self.table.rowCount())
            self.saved_state = self.snapshot()
            return True
        except Exception as exc:
            self.logger.exception('save_failed')
            QMessageBox.critical(self, '保存失败', f'{exc}\n请检查数据格式、文件占用或保存目录权限。')
            return False

if __name__ == '__main__':
    if '--self-test' in sys.argv:
        import json
        from diagnose_ths import collect_report
        app = QApplication([sys.argv[0], '-platform', 'offscreen'])
        win = MainWindow()
        report = {'window_created': True, 'ths': collect_report(),
                  'date_marking_verified': False, 'frozen': bool(getattr(sys, 'frozen', False))}
        position = sys.argv.index('--self-test')
        if len(sys.argv) > position + 1:
            Path(sys.argv[position + 1]).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps(report, ensure_ascii=True))
        win.close()
        sys.exit(0)
    app = QApplication(sys.argv); win = MainWindow(); win.show(); sys.exit(app.exec())
